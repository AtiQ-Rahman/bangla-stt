"""More output formats: WebVTT, Word (.docx), Bangla digits, and video with burned-in subtitles."""
import re
import subprocess
from pathlib import Path
from typing import Iterator

from stt.transcribe import Segment, _cues, _ts

BANGLA_DIGITS = str.maketrans("0123456789", "০১২৩৪৫৬৭৮৯")


def bangla_digits(text: str) -> str:
    """20 -> ২০."""
    return text.translate(BANGLA_DIGITS)


def to_vtt(segments: list[Segment]) -> str:
    """WebVTT subtitles (what YouTube Studio and HTML players take)."""
    cues = "\n".join(f"{_ts(start).replace(',', '.')} --> {_ts(end).replace(',', '.')}\n{text}\n"
                     for start, end, text in _cues(segments))
    return "WEBVTT\n\n" + cues


def srt_from_cues(cues: list[tuple[float, float, str]]) -> str:
    """SRT from (start, end, text) cues, e.g. translated ones."""
    return "\n".join(f"{i}\n{_ts(start)} --> {_ts(end)}\n{text}\n"
                     for i, (start, end, text) in enumerate(cues, 1))


def write_docx(text: str, path: Path, title: str) -> Path:
    """A Word document with the title as heading and one paragraph per transcript paragraph."""
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for para in re.split(r"\n\s*\n", text.strip()):
        doc.add_paragraph(para.strip())
    doc.save(path)
    return path


def _bangla_font() -> str:
    """A font family that has Bangla letters, for subtitles burned into video."""
    out = subprocess.run(["fc-match", "-f", "%{family[0]}", ":lang=bn"], capture_output=True, text=True)
    return out.stdout.strip() or "FreeSans"


def burn_subtitles(video: Path, srt: Path, out: Path) -> Iterator[float]:
    """Render the subtitles onto the video with ffmpeg, yielding progress 0..1 as it goes."""
    from stt.audio import duration

    total = duration(video) or 1.0
    style = f"FontName={_bangla_font()},FontSize=22,Outline=2,MarginV=28"
    # The subtitles filter parses its argument, so the path is passed relative to cwd.
    cmd = ["ffmpeg", "-nostdin", "-y", "-loglevel", "error", "-progress", "pipe:1", "-i", video.name,
           "-vf", f"subtitles={srt.name}:force_style='{style}'",
           "-c:v", "libx264", "-preset", "veryfast", "-crf", "23", "-c:a", "copy", str(out)]
    proc = subprocess.Popen(cmd, cwd=video.parent, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            text=True)
    try:
        for line in proc.stdout:
            if line.startswith("out_time_us="):
                us = line.split("=", 1)[1].strip()
                if us.isdigit():
                    yield min(1.0, int(us) / 1e6 / total)
    finally:
        if proc.poll() is None:  # stopped from the UI: don't leave ffmpeg running
            proc.kill()
        proc.wait()
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.read().strip().splitlines()[-1:] or "ffmpeg failed")
    yield 1.0
